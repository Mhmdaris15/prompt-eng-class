"""Local Markdown export to HTML and Word; no API calls or network access."""

from pathlib import Path


def export_document(path: Path) -> None:
    from markdown_it import MarkdownIt
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    text = path.read_text(encoding="utf-8")
    md = MarkdownIt("commonmark", {"html": True}).enable("table")
    css = """body{font:11pt/1.5 Georgia,serif;max-width:1100px;margin:36px auto;padding:0 24px;color:#17202a}h1,h2,h3{font-family:Arial,sans-serif}table{border-collapse:collapse;width:100%;font:9pt/1.3 Arial,sans-serif;margin:18px 0}td,th{border:1px solid #b8c2cc;padding:6px;vertical-align:top;overflow-wrap:anywhere}th{background:#edf2f7}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f5f7f9;padding:12px;font:9pt/1.35 Consolas,monospace}a{color:#185b97}@media print{body{margin:0;padding:0;max-width:none}h2,h3{break-after:avoid}tr{break-inside:avoid}thead{display:table-header-group}}@page{size:A4;margin:18mm}"""
    path.with_suffix(".html").write_text(
        '<!doctype html><html lang="ru"><meta charset="utf-8"><title>' + path.stem
        + '</title><style>' + css + '</style><body>' + md.render(text) + '</body></html>', encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.6)
    section.top_margin = section.bottom_margin = Cm(1.8)
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Times New Roman', Pt(11)
    style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    page = OxmlElement('w:fldSimple'); page.set(qn('w:instr'), 'PAGE'); footer._p.append(page)

    def inline(paragraph, token):
        bold = italic = False
        link = None
        for child in token.children or []:
            if child.type == 'strong_open': bold = True
            elif child.type == 'strong_close': bold = False
            elif child.type == 'em_open': italic = True
            elif child.type == 'em_close': italic = False
            elif child.type == 'link_open': link = child.attrGet('href')
            elif child.type == 'link_close': link = None
            elif child.type in ('softbreak', 'hardbreak'): paragraph.add_run('\n')
            elif child.type in ('text', 'code_inline'):
                if link:
                    node = OxmlElement('w:hyperlink')
                    node.set(qn('r:id'), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
                    run = OxmlElement('w:r'); properties = OxmlElement('w:rPr')
                    color = OxmlElement('w:color'); color.set(qn('w:val'), '185B97'); properties.append(color)
                    run.append(properties); content = OxmlElement('w:t'); content.text = child.content
                    run.append(content); node.append(run); paragraph._p.append(node)
                else:
                    run = paragraph.add_run(child.content); run.bold = bold; run.italic = italic
                    if child.type == 'code_inline': run.font.name, run.font.size = 'Consolas', Pt(9)

    tokens = md.parse(text)
    i, stack, prefix = 0, [], ''
    while i < len(tokens):
        token = tokens[i]
        if token.type in ('ordered_list_open', 'bullet_list_open'):
            stack.append([token.type, int(token.attrGet('start') or 1)])
        elif token.type in ('ordered_list_close', 'bullet_list_close'):
            stack.pop()
        elif token.type == 'list_item_open' and stack:
            prefix = f'{stack[-1][1]}. ' if stack[-1][0] == 'ordered_list_open' else '• '
            stack[-1][1] += 1
        if token.type == 'heading_open':
            if tokens[i + 1].content.startswith(('Приложение', '2. Цель')): doc.add_page_break()
            inline(doc.add_heading('', level=int(token.tag[1])), tokens[i + 1]); i += 3; continue
        if token.type == 'paragraph_open':
            inline(doc.add_paragraph(prefix), tokens[i + 1]); prefix = ''; i += 3; continue
        if token.type == 'fence':
            p = doc.add_paragraph()
            run = p.add_run(token.content.rstrip('\n')); run.font.name, run.font.size = 'Consolas', Pt(8)
        if token.type == 'table_open':
            matrix, current, j = [], [], i + 1
            while tokens[j].type != 'table_close':
                if tokens[j].type == 'tr_open': current = []
                elif tokens[j].type == 'inline': current.append(tokens[j])
                elif tokens[j].type == 'tr_close': matrix.append(current)
                j += 1
            table = doc.add_table(rows=0, cols=len(matrix[0])); table.style = 'Table Grid'
            for number, record in enumerate(matrix):
                cells = table.add_row().cells
                for cell, content in zip(cells, record):
                    inline(cell.paragraphs[0], content)
                    for p in cell.paragraphs:
                        for run in p.runs: run.font.size = Pt(8); run.bold = number == 0 or run.bold
                if number == 0:
                    header = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(header)
            i = j
        i += 1
    doc.save(path.with_suffix('.docx'))
